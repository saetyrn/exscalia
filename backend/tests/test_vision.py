import base64
import io

import pytest

from app.overseer.provider import image_block, to_gemini_contents, to_openai_messages
from app.overseer.vision import MAX_EDGE, attachments_in, is_viewable, load_image
from app.vault.notes import NoteError, NoteStore


def _png(width: int, height: int) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def _write_attachment(tmp_path, name: str, data: bytes) -> NoteStore:
    store = NoteStore(tmp_path)
    target = store.resolve_attachment(f"attachments/{name}")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return store


def test_small_image_passes_through_untouched(tmp_path):
    raw = _png(64, 48)
    store = _write_attachment(tmp_path, "diagram-a1.png", raw)

    loaded = load_image(store, "attachments/diagram-a1.png")

    assert loaded["media_type"] == "image/png"
    assert base64.b64decode(loaded["data"]) == raw


def test_large_image_is_downscaled_for_the_provider(tmp_path):
    from PIL import Image

    store = _write_attachment(tmp_path, "whiteboard-b2.png", _png(4000, 3000))

    loaded = load_image(store, "attachments/whiteboard-b2.png")

    # Re-encoded as JPEG and shrunk to the long-edge budget.
    assert loaded["media_type"] == "image/jpeg"
    with Image.open(io.BytesIO(base64.b64decode(loaded["data"]))) as shrunk:
        assert max(shrunk.size) <= MAX_EDGE


def test_non_images_and_missing_files_are_refused(tmp_path):
    store = _write_attachment(tmp_path, "slides-c3.pdf", b"%PDF-1.4 not an image")

    assert not is_viewable("attachments/slides-c3.pdf")
    with pytest.raises(NoteError):
        load_image(store, "attachments/slides-c3.pdf")
    with pytest.raises(NoteError):
        load_image(store, "attachments/never-existed.png")
    # Traversal is still rejected here, not just on the HTTP route.
    with pytest.raises(NoteError):
        load_image(store, "attachments/../../etc/passwd.png")


def test_attachment_links_are_found_in_note_bodies():
    body = (
        "# Lecture\n"
        "![board](/api/attachments/whiteboard-b2.png)\n"
        "[slides](/api/attachments/slides-c3.pdf)\n"
        "![again](/api/attachments/whiteboard-b2.png)\n"
        "[external](https://example.com/not-an-attachment.png)\n"
    )
    assert attachments_in(body) == [
        "attachments/whiteboard-b2.png",
        "attachments/slides-c3.pdf",
    ]


def test_image_blocks_translate_for_every_provider():
    block = image_block("image/png", "QUJD")
    history = [{"role": "user", "content": [{"type": "text", "text": "what is this?"}, block]}]

    # Anthropic consumes the neutral shape directly.
    assert block["source"]["media_type"] == "image/png"

    openai = to_openai_messages(history)
    parts = openai[0]["content"]
    assert parts[0]["type"] == "text"
    assert parts[1]["image_url"]["url"] == "data:image/png;base64,QUJD"

    gemini = to_gemini_contents(history)
    inline = gemini[0]["parts"][1]["inline_data"]
    assert inline == {"mime_type": "image/png", "data": "QUJD"}


def test_text_only_turns_keep_their_simple_shape():
    history = [{"role": "user", "content": [{"type": "text", "text": "hello"}]}]
    assert to_openai_messages(history) == [{"role": "user", "content": "hello"}]


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for line in paragraphs:
        doc.add_paragraph(line)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_text_is_extracted_with_paragraphs_and_tables(tmp_path):
    from app.overseer.vision import is_text_document, load_text_document

    raw = _docx(
        ["Meeting notes", "Ship the docx feature."],
        table=[["Task", "Owner"], ["Backend", "Ada"]],
    )
    store = _write_attachment(tmp_path, "notes-d4.docx", raw)

    assert is_text_document("attachments/notes-d4.docx")
    loaded = load_text_document(store, "attachments/notes-d4.docx")

    assert loaded["name"] == "notes-d4.docx"
    assert "Meeting notes" in loaded["text"]
    assert "Ship the docx feature." in loaded["text"]
    # Table cells are flattened row by row.
    assert "Task | Owner" in loaded["text"]
    assert "Backend | Ada" in loaded["text"]


def test_docx_is_truncated_when_oversized(tmp_path):
    from app.overseer.vision import MAX_TEXT_DOCUMENT_CHARS, load_text_document

    big = _docx([("lorem ipsum " * 40) for _ in range(400)])
    store = _write_attachment(tmp_path, "huge-e5.docx", big)

    loaded = load_text_document(store, "attachments/huge-e5.docx")

    assert "truncated" in loaded["text"]
    assert len(loaded["text"]) <= MAX_TEXT_DOCUMENT_CHARS + 200


def test_docx_rejects_non_docx_corrupt_and_missing(tmp_path):
    from app.overseer.vision import load_text_document

    # Wrong type entirely.
    png_store = _write_attachment(tmp_path, "pic-f6.png", _png(8, 8))
    with pytest.raises(NoteError):
        load_text_document(png_store, "attachments/pic-f6.png")

    # A .docx extension over bytes that are not a real zip/docx.
    bad_store = _write_attachment(tmp_path, "broken-g7.docx", b"not a real docx")
    with pytest.raises(NoteError) as exc:
        load_text_document(bad_store, "attachments/broken-g7.docx")
    assert exc.value.status == 400

    # Missing file.
    with pytest.raises(NoteError):
        load_text_document(bad_store, "attachments/ghost-h8.docx")

    # Traversal is refused here too.
    with pytest.raises(NoteError):
        load_text_document(bad_store, "attachments/../../etc/passwd.docx")
