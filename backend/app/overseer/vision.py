"""Turn vault attachments into image blocks the models can actually see.

Notes can hold pasted slides and whiteboard photos; without this OMNIA only
ever sees the literal markdown `![alt](/api/attachments/photo-a1b2.png)` and
has no idea whether that is a lecture slide or a cat.
"""

from __future__ import annotations

import base64
import hashlib
import io
from collections import OrderedDict
from pathlib import Path

from ..vault import fm
from ..vault.notes import NoteError, NoteStore

# Anthropic's guidance is ~1568px on the long edge; beyond that costs tokens
# without adding legibility. The same size is comfortable for OpenAI and Gemini.
MAX_EDGE = 1568
# Hard ceiling on what we will hand a provider, after downscaling.
MAX_ENCODED_BYTES = 4 * 1024 * 1024

RASTER_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

DOCUMENT_TYPES = {".pdf": "application/pdf"}
# A slide deck is the realistic case. Beyond this the request itself gets
# unwieldy and the provider starts rejecting it.
MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_DOCUMENT_PAGES = 50

# Formats the providers can't take natively, so we extract their text and hand
# it to the model as a plain text block instead of a document/image block.
TEXT_DOCUMENT_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
# .docx is a zip of XML — the on-disk file is small but a malicious one can
# expand hugely, so cap the input bytes as well as the extracted text.
MAX_TEXT_DOCUMENT_BYTES = 8 * 1024 * 1024
# Roughly the whole PDF budget; past this the request bloats without the model
# gaining much. We cut with a visible note so nobody thinks it saw the rest.
MAX_TEXT_DOCUMENT_CHARS = 40_000


def is_viewable(rel_path: str) -> bool:
    return Path(rel_path).suffix.lower() in RASTER_TYPES


def is_document(rel_path: str) -> bool:
    return Path(rel_path).suffix.lower() in DOCUMENT_TYPES


def is_text_document(rel_path: str) -> bool:
    """A document we extract to text (e.g. .docx) rather than send natively."""
    return Path(rel_path).suffix.lower() in TEXT_DOCUMENT_TYPES


def is_attachable(rel_path: str) -> bool:
    """Anything OMNIA can be handed directly: images, PDFs and .docx."""
    return is_viewable(rel_path) or is_document(rel_path) or is_text_document(rel_path)


def load_text_document(store: NoteStore, rel_path: str) -> dict[str, str]:
    """Extract plain text from a .docx so a provider that can't take it natively
    still gets the content as a text block.

    Paragraphs (headings and list items included, as plain text) are joined by
    newlines; table cells are flattened row by row. Images, drawings and exact
    formatting are dropped — this is a text downgrade, and the caller frames it
    as such. A corrupt or unreadable file raises NoteError(400).
    """
    if not is_text_document(rel_path):
        raise NoteError(400, f"{rel_path} is not a text document")
    abs_path = store.resolve_attachment(rel_path)
    if not abs_path.is_file():
        raise NoteError(404, f"Attachment not found: {rel_path}")

    size = abs_path.stat().st_size
    if size > MAX_TEXT_DOCUMENT_BYTES:
        raise NoteError(
            413,
            f"{rel_path} is {size // (1024 * 1024)} MB — too large to read. "
            f"The limit is {MAX_TEXT_DOCUMENT_BYTES // (1024 * 1024)} MB.",
        )

    try:
        from docx import Document

        document = Document(str(abs_path))
    except Exception as exc:  # python-docx raises many types on a bad zip/xml
        raise NoteError(
            400,
            f"Could not read {abs_path.name} — it may be corrupt or not a real "
            f".docx file.",
        ) from exc

    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)

    text = "\n".join(parts).strip()
    if len(text) > MAX_TEXT_DOCUMENT_CHARS:
        text = (
            text[:MAX_TEXT_DOCUMENT_CHARS]
            + "\n\n…(truncated — the document was longer than "
            f"{MAX_TEXT_DOCUMENT_CHARS} characters)"
        )
    return {"name": abs_path.name, "text": text}


def load_document(store: NoteStore, rel_path: str) -> dict[str, str]:
    """Read a PDF as base64 for providers that accept documents natively."""
    if not is_document(rel_path):
        raise NoteError(400, f"{rel_path} is not a PDF")
    abs_path = store.resolve_attachment(rel_path)
    if not abs_path.is_file():
        raise NoteError(404, f"Attachment not found: {rel_path}")

    # Check the size before reading, not after — a 2 GB file used to be pulled
    # into memory in full just to be rejected on the next line.
    size = abs_path.stat().st_size
    if size > MAX_DOCUMENT_BYTES:
        raise NoteError(
            413,
            f"{rel_path} is {size // (1024 * 1024)} MB — too large to send. "
            f"The limit is {MAX_DOCUMENT_BYTES // (1024 * 1024)} MB.",
        )
    raw = abs_path.read_bytes()
    pages = page_count(raw)
    if pages and pages > MAX_DOCUMENT_PAGES:
        raise NoteError(
            413,
            f"{rel_path} has {pages} pages — the limit is {MAX_DOCUMENT_PAGES}. "
            "Split it, or ask about a specific section.",
        )
    return {
        "media_type": DOCUMENT_TYPES[abs_path.suffix.lower()],
        "data": base64.b64encode(raw).decode("ascii"),
        "name": abs_path.name,
    }


def page_count(raw: bytes) -> int:
    try:
        from pypdf import PdfReader

        return len(PdfReader(io.BytesIO(raw)).pages)
    except Exception:
        return 0


# Extraction is synchronous and blocks the event loop, and the message list is
# rebuilt on every one of up to 15 agent rounds — so the same 50-page PDF was
# parsed 15 times per turn. Hashing the bytes is orders of magnitude cheaper.
_PDF_TEXT_CACHE: OrderedDict[str, str] = OrderedDict()
_PDF_TEXT_CACHE_MAX = 8


def extract_pdf_text(raw: bytes, limit: int = 60_000) -> str:
    """Plain text from a PDF, for providers that cannot take documents.

    Loses layout, figures and handwriting — a real downgrade, so callers say so
    rather than pretending the model saw the document.
    """
    key = f"{hashlib.sha256(raw).hexdigest()}:{limit}"
    if key in _PDF_TEXT_CACHE:
        _PDF_TEXT_CACHE.move_to_end(key)
        return _PDF_TEXT_CACHE[key]
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        parts = []
        for page in reader.pages[:MAX_DOCUMENT_PAGES]:
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) > limit:
                break
        text = "\n\n".join(parts).strip()
    except Exception:
        return ""
    result = text[:limit]
    _PDF_TEXT_CACHE[key] = result
    while len(_PDF_TEXT_CACHE) > _PDF_TEXT_CACHE_MAX:
        _PDF_TEXT_CACHE.popitem(last=False)
    return result


def load_image(store: NoteStore, rel_path: str) -> dict[str, str]:
    """Read an attachment and return `{media_type, data}` (base64), downscaled
    so a phone photo doesn't blow the request size."""
    if not is_viewable(rel_path):
        raise NoteError(
            400,
            f"{rel_path} is not a viewable image. Supported: "
            + ", ".join(sorted(RASTER_TYPES)),
        )
    abs_path = store.resolve_attachment(rel_path)
    if not abs_path.is_file():
        raise NoteError(404, f"Attachment not found: {rel_path}")

    raw = abs_path.read_bytes()
    return encode_image(raw, RASTER_TYPES[abs_path.suffix.lower()], label=rel_path)


def encode_image(
    raw: bytes, media_type: str, *, label: str = "The image"
) -> dict[str, str]:
    """Downscale raw image bytes and base64 them into a `{media_type, data}`
    block ready for :func:`provider.image_block`.

    Factored out of :func:`load_image` because not every image OMNIA looks at
    is an attachment on disk — a Paper page render is bytes inside a notebook
    package, and it deserves the same size discipline.
    """
    data, media_type = _downscale(raw, media_type)
    if len(data) > MAX_ENCODED_BYTES:
        raise NoteError(413, f"{label} is too large to send even after resizing")
    return {"media_type": media_type, "data": base64.b64encode(data).decode("ascii")}


def _downscale(raw: bytes, media_type: str) -> tuple[bytes, str]:
    """Shrink to MAX_EDGE if needed. Returns the original bytes untouched when
    Pillow can't read it — better a large image than none."""
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover - Pillow is a hard dependency
        return raw, media_type
    try:
        with Image.open(io.BytesIO(raw)) as image:
            if max(image.size) <= MAX_EDGE and len(raw) <= MAX_ENCODED_BYTES:
                return raw, media_type
            image = image.convert("RGB")
            image.thumbnail((MAX_EDGE, MAX_EDGE))
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=85, optimize=True)
            return buffer.getvalue(), "image/jpeg"
    except Exception:
        return raw, media_type


def attachments_in(content: str) -> list[str]:
    """Vault-relative attachment paths referenced by a note, frontmatter first."""
    import re

    # Links are written as /api/attachments/<name> by the editor.
    marker = "/api/attachments/"
    found: list[str] = []
    if marker not in content:
        return found
    # Media art (cover:, backdrop:) is a plain frontmatter value, not a body
    # link — missing it reported every live cover as an orphan to delete.
    try:
        meta, _body = fm.parse(content)
    except Exception:
        meta = {}
    for value in meta.values():
        if isinstance(value, str) and value.strip().startswith(marker):
            rel = f"{NoteStore.ATTACHMENT_DIR}/{value.strip()[len(marker):]}"
            if rel not in found:
                found.append(rel)
    for match in re.finditer(r"!?\[[^\]]*\]\(([^)\s]+)\)", content):
        url = match.group(1)
        # Links are written as /api/attachments/<name> by the editor.
        if marker in url:
            name = url.split(marker, 1)[1]
            rel = f"{NoteStore.ATTACHMENT_DIR}/{name}"
            if rel not in found:
                found.append(rel)
    return found
